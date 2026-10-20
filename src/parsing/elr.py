"""Lecture de l'Equipment List Report (ELR) au format Excel livré par iBwave.

L'ELR a une ligne par modèle : type, fabricant, modèle, description,
numéro d'inventaire et quantité. La quantité est un nombre, sauf pour les
câbles où elle est écrite avec son unité (« 1245,22 m »).
"""

import openpyxl
import pandas as pd

from .utils import clean, split_unit

REQUIRED = {"Type", "Model", "Qty"}


def _find_header(ws):
    """Renvoie (numéro de ligne, index de colonne par nom) de l'en-tête de l'ELR."""
    for row in ws.iter_rows():
        values = [clean(c.value) for c in row]
        if REQUIRED.issubset(values):
            return row[0].row, {name: values.index(name) for name in values if name}
    raise ValueError("En-tête ELR introuvable (colonnes Type, Model, Qty attendues).")


def parse_elr(path):
    """Renvoie un DataFrame : une ligne par modèle du bordereau."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header_row, cols = _find_header(ws)

    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        first = clean(row[0]) if row else ""
        if first.lower().startswith("created on"):
            break
        model = clean(row[cols["Model"]]) if cols.get("Model") is not None else ""
        if not model:
            continue
        qty_raw = row[cols["Qty"]]
        qty, unit = split_unit(qty_raw)
        rows.append(
            {
                "type": clean(row[cols["Type"]]) if cols.get("Type") is not None else "",
                "manufacturer": clean(row[cols["Manufacturer"]]) if cols.get("Manufacturer") is not None else "",
                "model": model,
                "description": clean(row[cols["Description"]]) if cols.get("Description") is not None else "",
                "inventory": clean(row[cols["Inventory#"]]) if cols.get("Inventory#") is not None else "",
                "qty": qty,
                "unit": unit,
            }
        )
    return pd.DataFrame(rows, columns=["type", "manufacturer", "model", "description", "inventory", "qty", "unit"])